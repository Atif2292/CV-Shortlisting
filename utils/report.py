"""
utils/report.py
Excel shortlist report with a UK GDPR audit sheet that records what actually happened.
"""

import datetime
import io
from typing import Any, Dict, List

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from utils import config

BANDS = ("strong", "good", "moderate", "weak", "invalid")
BAND_LABELS = {
    "strong": "Strong Match",
    "good": "Good Match",
    "moderate": "Review",
    "weak": "Weak Match",
    "invalid": "Not a CV?",
}
MODE_LABELS = {
    "cloud":   "Cloud AI (OpenAI)",
    "offline": "Private Engine (no external AI calls)",
    "mixed":   "Cloud AI, with Private Engine fallback for some CVs",
}


def score(r: Dict[str, Any]) -> int:
    try:
        return max(0, min(100, int(round(float(r.get("match_score", 0))))))
    except (TypeError, ValueError):
        return 0


def fit_band(r: Dict[str, Any]) -> str:
    band = r.get("fit_band")
    if band in BANDS:
        return band
    s = score(r)
    return "strong" if s >= 80 else "good" if s >= 65 else "moderate" if s >= 45 else "weak"


def xl_safe(value):
    """Stop spreadsheet formula injection from CV or AI text."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def _style(ws, header_color: str, row_height: int, widths: List[int], valign: str) -> None:
    hfont = Font(bold=True, color="FFFFFF", size=11, name="Calibri")
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=header_color)
        cell.font = hfont
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28
    even, odd = PatternFill("solid", fgColor="EFF6FF"), PatternFill("solid", fgColor="FFFFFF")
    body = Font(size=10, name="Calibri", color="0B1120")
    side = Side(style="thin", color="E2E8F0")
    border = Border(left=side, right=side, top=side, bottom=side)
    for ri, row in enumerate(ws.iter_rows(min_row=2), start=2):
        for cell in row:
            cell.fill = even if ri % 2 == 0 else odd
            cell.font = body
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical=valign)
        ws.row_dimensions[ri].height = row_height
    for idx, width in enumerate(widths):
        ws.column_dimensions[chr(ord("A") + idx)].width = width


def build_excel_report(results: List[Dict[str, Any]], overrides: Dict[Any, Dict[str, Any]],
                       run_meta: Dict[str, Any], audit_stats: Dict[str, Any], controller: str) -> bytes:
    reviewed = sum(1 for o in overrides.values() if o.get("decided_at"))
    rows = []
    for i, r in enumerate(results):
        ov = overrides.get(r.get("candidate_index", i), {})
        ai_score = score(r)
        delta = int(ov.get("score_delta", 0) or 0)
        rows.append({
            "Rank":                       i + 1,
            "Name":                       r.get("candidate_name", ""),
            "Email":                      r.get("email", ""),
            "AI Match Score":             f"{ai_score}%",
            "AI Fit":                     BAND_LABELS[fit_band(r)],
            "Recruiter Decision":         ov.get("status", "Under Review"),
            "Decision Recorded (UTC)":    ov.get("decided_at") or "Not reviewed",
            "Recruiter Score Adjustment": delta,
            "Final Score":                f"{max(0, min(100, ai_score + delta))}%",
            "Recruiter Notes":            ov.get("notes", ""),
            "Scoring Engine":             r.get("processing_mode", ""),
            "Strengths":                  "; ".join(str(s) for s in (r.get("strengths") or [])),
            "Concerns":                   "; ".join(str(c) for c in (r.get("concerns") or [])),
            "AI Recommendation":          r.get("recommendation", ""),
        })
    df_results = pd.DataFrame(rows).map(xl_safe)

    used = run_meta.get("mode_used", "")
    cloud_used = used in ("cloud", "mixed")
    redaction = (
        "Automated redaction applied before Cloud AI scoring (best-effort): "
        f"emails {audit_stats.get('emails', 0)}, phones {audit_stats.get('phones', 0)}, "
        f"postcodes {audit_stats.get('postcodes', 0)}, names {audit_stats.get('names', 0)}, "
        f"profile links {audit_stats.get('profiles', 0)}, dates of birth/ages {audit_stats.get('dob_age', 0)}, "
        f"demographic fields {audit_stats.get('demographics', 0)}."
        if cloud_used else "Not applicable — no CV content was sent to an external AI provider."
    )
    engine = MODE_LABELS.get(used, used or "Unknown")
    if run_meta.get("fallback_reason"):
        engine += f" — fallback reason: {run_meta['fallback_reason']}"
    audit_rows = [
        ("Statutory Framework", "UK GDPR & Data Protection Act 2018"),
        ("Governing Jurisdiction", "England and Wales"),
        ("Data Controller", controller or "Customer"),
        ("Data Processor", f"{config.COMPANY_NAME} (RankCV), under the RankCV Data Processing Agreement"),
        ("Run ID / Time", f"{run_meta.get('run_id', '')} · {run_meta.get('started_at', '')}"),
        ("Scoring Engine Used", engine),
        ("External AI Sub-processor",
         "OpenAI (United States) — redacted CV text only" if cloud_used else "None (Private Engine)"),
        ("Article 22 Human Oversight",
         f"{reviewed} of {len(results)} candidates have a decision recorded by a recruiter in RankCV. "
         "Candidates still 'Under Review' have no human decision and must not be rejected on the AI score alone."),
        ("Article 5(1)(c) Data Minimisation", redaction),
        ("Article 5(1)(e) Storage Limitation",
         "CV files were held only in a per-run temporary folder and deleted once text was extracted. "
         "Results are kept in the browser session's memory only and are not stored in RankCV's database."),
        ("Equality Act 2010",
         "Scoring compares skills, experience and qualifications with the job description. "
         "Common protected-characteristic fields are redacted before Cloud AI scoring (best-effort). "
         "Recruiters should check results for possible bias before deciding."),
        ("Report Generated (UTC)", datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")),
        ("Recruiter Sign-Off",
         "Name: ____________________   Date: ____________   (complete before relying on this shortlist)"),
    ]
    df_audit = pd.DataFrame(audit_rows, columns=["Audit Parameter", "Record"]).map(xl_safe)

    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        df_results.to_excel(xw, index=False, sheet_name="RankCV Results")
        ws = xw.sheets["RankCV Results"]
        _style(ws, "2563EB", 52, [7, 22, 28, 12, 14, 18, 20, 12, 12, 30, 28, 36, 36, 50], "top")
        ws.freeze_panes = "A2"
        df_audit.to_excel(xw, index=False, sheet_name="UK GDPR Audit")
        _style(xw.sheets["UK GDPR Audit"], "0F766E", 40, [34, 95], "center")
    return buf.getvalue()
